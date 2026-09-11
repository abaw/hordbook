"""Readers for the official NGSL 1.2 downloads (CC BY-SA 4.0).

Two files are used:
- ``NGSL_12_stats.csv``: lemma and SFI frequency rank.
- ``NGSL_12_with_English_definitions.xlsx``: lemma and easy-English definition.

Both are normalised to a canonical lemma form so they can be joined 1:1.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

import openpyxl

# Spellings that differ between the two official files. The lemmatized
# teaching list is treated as canonical (it spells "email" without a hyphen).
LEMMA_ALIASES: dict[str, str] = {
    "e-mail": "email",
}


def canonical_lemma(raw: str) -> str:
    """Normalise a lemma as it appears in either official file.

    - Excel turned ``true``/``false`` into the booleans ``TRUE``/``FALSE`` in
      the stats CSV; lowercasing restores them. The only legitimately
      capitalised lemma, ``I``, is preserved.
    - Known spelling differences are mapped through ``LEMMA_ALIASES``.
    """
    lemma = raw.strip()
    if lemma != "I":
        lemma = lemma.lower()
    return LEMMA_ALIASES.get(lemma, lemma)


@dataclass(frozen=True)
class RankedLemma:
    lemma: str
    rank: int


def read_stats_csv(path: Path) -> list[RankedLemma]:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        return [
            RankedLemma(lemma=canonical_lemma(row["Lemma"]), rank=int(row["SFI Rank"]))
            for row in reader
            if row.get("Lemma")
        ]


def read_definitions_xlsx(path: Path) -> dict[str, str]:
    """Return ``{lemma: definition}`` from the official definitions workbook.

    The workbook has exactly two columns (word, definition); the header's
    spelling is not relied upon.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    next(rows, None)  # header
    definitions: dict[str, str] = {}
    for row in rows:
        if not row or row[0] is None:
            continue
        lemma = canonical_lemma(str(row[0]))
        definition = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
        definitions[lemma] = definition
    wb.close()
    return definitions
